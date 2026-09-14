"""把拣货效率计算结果写入 门店实时效率监控报表V1.0.xlsx 的新工作表。
采用 zip 部件级改写：仅新增 worksheet 部件并更新 workbook.xml / rels / [Content_Types].xml，
不改动 Power Query、透视表、图片等既有部件，避免 openpyxl 保存造成的破坏。
"""
import datetime as dt
import os
import re
import shutil
import zipfile
from collections import OrderedDict

import openpyxl

REPORT = r'h:\Hema-Picking-Efficiency\Hema-Picking-Efficiency\门店实时效率监控报表V1.0.xlsx'
BACKUP = r'h:\Hema-Picking-Efficiency\Hema-Picking-Efficiency\门店实时效率监控报表V1.0_备份.xlsx'
EXPORT = r'h:\Hema-Picking-Efficiency\Hema-Picking-Efficiency\export-1785428943043.xlsx'

# ---------- 1. 计算 ----------
mapping = {}
for r in openpyxl.load_workbook(REPORT, data_only=True)['前后场填写'].iter_rows(min_row=2, values_only=True):
    if r and r[0] is not None:
        mapping[str(r[0]).strip()] = str(r[1]).strip() if r[1] else ''

ws = openpyxl.load_workbook(EXPORT, data_only=True)['data']
data = list(ws.iter_rows(values_only=True))
hdr = list(data[0])
ix = {}
for i, h in enumerate(hdr):
    if str(h).strip() not in ix:
        ix[str(h).strip()] = i


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def as_dt(v):
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, str) and v.strip():
        for f in ('%Y-%m-%d %H:%M:%S', '%Y/%m/%d %H:%M:%S'):
            try:
                return dt.datetime.strptime(v.strip(), f)
            except ValueError:
                pass
    return None


def job_type(part, sub):
    if sub == '拆零拣打一体':
        return '一体化'
    p = mapping.get(str(part).strip())
    return p or '未匹配分区'


def zone_code(z):
    s = str(z or '').strip()
    i = s.find(' ')
    return s[:i] if i > 0 else s


recs = []
for r in data[1:]:
    if r[ix['拣货单号']] is None or not r[ix['拣货人']]:
        continue
    t0, t1 = as_dt(r[ix['拣货开始时间']]), as_dt(r[ix['拣货完成时间']])
    if not t0 or not t1:
        continue
    hrs = (t1 - t0).total_seconds() / 3600.0
    if hrs <= 0:
        continue
    zone = str(r[ix['拣货分区']] or '').strip()
    recs.append({'小时': t0.hour,
                 '作业类型': job_type(zone, r[ix['任务子类型']]),
                 '拣货分区': zone, '分区简码': zone_code(zone),
                 '拣货人': str(r[ix['拣货人']]).strip(),
                 '行数': num(r[ix['拣货行数']]), '时长': hrs})


def agg(rs):
    """返回 (拣货行数合计, 时长合计h) 原始值，不取整"""
    return sum(x['行数'] for x in rs), sum(x['时长'] for x in rs)


def eff(h, t):
    return round(h / t, 2) if t else None


def group(keyfunc):
    g = OrderedDict()
    for x in recs:
        g.setdefault(keyfunc(x), []).append(x)
    return g


hours = sorted({x['小时'] for x in recs})
H, T = agg(recs)
E = eff(H, T)

# 表A 按作业类型
sheetA = [['作业类型', '拣货行数', '时长(h)', '效率(行/h)']]
gA = group(lambda x: x['作业类型'])
for k in sorted(gA, key=lambda k: -eff(*agg(gA[k]))):
    h, t = agg(gA[k])
    sheetA.append([k, round(h), round(t, 4), eff(h, t)])
sheetA.append(['总计', round(H), round(T, 4), E])

# 表B 按人
sheetB = [['拣货人', '拣货行数', '时长(h)', '效率(行/h)']]
gB = group(lambda x: x['拣货人'])
for k in sorted(gB, key=lambda k: -eff(*agg(gB[k]))):
    h, t = agg(gB[k])
    sheetB.append([k, round(h), round(t, 4), eff(h, t)])
sheetB.append(['总计', round(H), round(T, 4), E])

# 表C 按小时（升序）
sheetC = [['小时', '拣货行数', '时长(h)', '效率(行/h)']]
gh = group(lambda x: x['小时'])
for h in hours:
    a, b = agg(gh[h])
    sheetC.append(['%d点' % h, round(a), round(b, 4), eff(a, b)])
sheetC.append(['总计', round(H), round(T, 4), E])

# 表D 作业类型 x 小时
sheetD = [['作业类型'] + ['%d点' % h for h in hours] + ['全天']]
g2 = group(lambda x: (x['作业类型'], x['小时']))
for j in sorted(gA, key=lambda k: -eff(*agg(gA[k]))):
    row = [j]
    rh = rt = 0.0
    for h in hours:
        if (j, h) in g2:
            a, b = agg(g2[(j, h)])
            rh += a
            rt += b
            row.append(eff(a, b))
        else:
            row.append(None)
    row.append(eff(rh, rt))
    sheetD.append(row)
sheetD.append(['合计'] + [eff(*agg(gh[h])) for h in hours] + [E])

# 表E 作业类型+拣货人 x 小时
sheetE = [['作业类型', '拣货人'] + ['%d点' % h for h in hours] + ['总计']]
g3 = group(lambda x: (x['作业类型'], x['拣货人']))
for jt, p in sorted(g3.keys(), key=lambda k: -eff(*agg(g3[k]))):
    row = [jt, p]
    th = tt = 0.0
    for h in hours:
        rs = [x for x in g3[(jt, p)] if x['小时'] == h]
        if rs:
            a, b = agg(rs)
            th += a
            tt += b
            row.append(eff(a, b))
        else:
            row.append(None)
    row.append(eff(th, tt))
    sheetE.append(row)
sheetE.append(['合计', ''] + [eff(*agg(gh[h])) for h in hours] + [E])

# 表F 细分：作业类型 × 拣货分区
type_order = sorted(gA, key=lambda k: -eff(*agg(gA[k])))
gZ = group(lambda x: (x['作业类型'], x['拣货分区']))
sheetZ = [['作业类型', '拣货分区', '拣货行数', '时长(h)', '效率(行/h)', '行数占比(%)']]
zrows = []
for jt in type_order:
    items = [(z, rs) for (t, z), rs in gZ.items() if t == jt]
    for z, rs in sorted(items, key=lambda kv: -eff(*agg(kv[1]))):
        a, b = agg(rs)
        zrows.append((jt, z, rs))
        sheetZ.append([jt, z, round(a), round(b, 4), eff(a, b), round(a / (H or 1) * 1000) / 10])
sheetZ.append(['合计', '', round(H), round(T, 4), E, 100.0])

# 表G 细分 × 小时
sheetZH = [['作业类型', '拣货分区'] + ['%d点' % h for h in hours] + ['全天']]
for jt, z, rs in zrows:
    row = [jt, z]
    th = tt = 0.0
    for h in hours:
        sub = [x for x in rs if x['小时'] == h]
        if sub:
            a, b = agg(sub)
            th += a
            tt += b
            row.append(eff(a, b))
        else:
            row.append(None)
    row.append(eff(th, tt))
    sheetZH.append(row)
sheetZH.append(['合计', ''] + [eff(*agg(gh[h])) for h in hours] + [E])

SHEETS = [('效率-按作业类型', sheetA), ('效率-按小时', sheetC),
          ('效率-作业类型x小时', sheetD), ('效率-作业类型拣货人x小时', sheetE),
          ('效率-按拣货分区(细分)', sheetZ), ('效率-细分x小时', sheetZH),
          ('效率-按人', sheetB)]

# ---------- 2. 生成 worksheet XML（内联字符串，无需 sharedStrings） ----------
def colref(n):
    s = ''
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def esc(s):
    return (str(s).replace('&', '&amp;').replace('<', '&lt;')
            .replace('>', '&gt;').replace('"', '&quot;'))


def sheet_xml(rows):
    parts = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
             '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>']
    for ri, row in enumerate(rows, 1):
        cells = []
        for ci, v in enumerate(row, 1):
            if v is None or v == '':
                continue
            ref = colref(ci) + str(ri)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cells.append('<c r="%s"><v>%s</v></c>' % (ref, v))
            else:
                cells.append('<c r="%s" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (ref, esc(v)))
        if cells:
            parts.append('<row r="%d">%s</row>' % (ri, ''.join(cells)))
    parts.append('</sheetData></worksheet>')
    return ''.join(parts).encode('utf-8')


# ---------- 3. zip 部件级改写 ----------
# 以首次运行时的原始文件作为只读源，结果写入 REPORT，可反复执行不叠加
if not os.path.exists(BACKUP):
    shutil.copyfile(REPORT, BACKUP)
zin = zipfile.ZipFile(BACKUP)
orig = {n: zin.read(n) for n in zin.namelist()}
zin.close()

wbxml = orig['xl/workbook.xml'].decode('utf-8')
rels = orig['xl/_rels/workbook.xml.rels'].decode('utf-8')
ctypes = orig['[Content_Types].xml'].decode('utf-8')

max_sheet = max(int(m) for m in re.findall(r'worksheets/sheet(\d+)\.xml', ' '.join(orig.keys())))
max_sheetid = max(int(m) for m in re.findall(r'sheetId="(\d+)"', wbxml))
max_rid = max(int(m) for m in re.findall(r'Id="rId(\d+)"', rels))

extra = {}
new_sheets_xml, new_rels, new_ctypes = [], [], []
for i, (name, rows) in enumerate(SHEETS, 1):
    part = 'xl/worksheets/sheet%d.xml' % (max_sheet + i)
    arc = 'worksheets/sheet%d.xml' % (max_sheet + i)
    rid = 'rId%d' % (max_rid + i)
    sid = max_sheetid + i
    extra[part] = sheet_xml(rows)
    new_sheets_xml.append('<sheet name="%s" sheetId="%d" r:id="%s"/>' % (esc(name), sid, rid))
    new_rels.append('<Relationship Id="%s" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    'relationships/worksheet" Target="%s"/>' % (rid, arc))
    new_ctypes.append('<Override PartName="/%s" ContentType="application/vnd.openxmlformats-officedocument.'
                      'spreadsheetml.worksheet+xml"/>' % part)

extra['xl/workbook.xml'] = wbxml.replace('</sheets>', ''.join(new_sheets_xml) + '</sheets>').encode('utf-8')
extra['xl/_rels/workbook.xml.rels'] = rels.replace('</Relationships>', ''.join(new_rels) + '</Relationships>').encode('utf-8')
extra['[Content_Types].xml'] = ctypes.replace('</Types>', ''.join(new_ctypes) + '</Types>').encode('utf-8')

with zipfile.ZipFile(REPORT, 'w', zipfile.ZIP_DEFLATED) as zout:
    for n in orig:
        zout.writestr(n, extra.get(n, orig[n]))
    for n, d in extra.items():
        if n not in orig:
            zout.writestr(n, d)

print('已写入工作表:', [s[0] for s in SHEETS])
print('总行数:', len(recs), '| 小时:', hours)
print('备份:', BACKUP)
