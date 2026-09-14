"""用 export-1785428943043.xlsx（原始拣货单导出）按报表口径计算拣货效率
效率达成 = Σ拣货行数 / Σ拣货时长(h)
作业类型 = IF(任务子类型="拆零拣打一体","一体化", VLOOKUP(拣货分区,前后场备注,2,0))
"""
import datetime as dt
from collections import defaultdict, OrderedDict

import openpyxl

TEMPLATE = r'h:\Hema-Picking-Efficiency\Hema-Picking-Efficiency\门店实时效率监控报表V1.0.xlsx'
EXPORT = r'h:\Hema-Picking-Efficiency\Hema-Picking-Efficiency\export-1785428943043.xlsx'

# 1) 前后场映射（沿用报表里的配置）
wbm = openpyxl.load_workbook(TEMPLATE, data_only=True)
mapping = {}
for r in wbm['前后场填写'].iter_rows(min_row=2, values_only=True):
    if r and r[0] is not None:
        mapping[str(r[0]).strip()] = str(r[1]).strip() if r[1] else ''

# 2) 导出明细
ws = openpyxl.load_workbook(EXPORT, data_only=True)['data']
data = list(ws.iter_rows(values_only=True))
hdr = list(data[0])
ix = {}
for i, h in enumerate(hdr):
    if h not in ix:          # 存在两个"任务类型"同名列，取第一个
        ix[str(h).strip()] = i


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def as_dt(v):
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, str) and v.strip():
        for f in ('%Y-%m-%d %H:%M:%S', '%Y/%m/%d %H:%M:%S'):
            try:
                return dt.datetime.strptime(v.strip(), f)
            except ValueError:
                pass
    return None


def job_type(part, sub):
    if sub == '拆零拣打一体':
        return '一体化'
    p = mapping.get(str(part).strip())
    return p or '未匹配分区'


def zone_code(z):
    s = str(z or '').strip()
    i = s.find(' ')
    return s[:i] if i > 0 else s


recs, unmatched = [], defaultdict(int)
skip_empty_person = skip_no_time = 0
for r in data[1:]:
    if r[ix['拣货单号']] is None:
        continue
    person = r[ix['拣货人']]
    if person is None or str(person).strip() == '':
        skip_empty_person += 1
        continue
    t0, t1 = as_dt(r[ix['拣货开始时间']]), as_dt(r[ix['拣货完成时间']])
    if not t0 or not t1:
        skip_no_time += 1
        continue
    hours = (t1 - t0).total_seconds() / 3600.0
    if hours <= 0:
        skip_no_time += 1
        continue
    part = r[ix['拣货分区']]
    jt = job_type(part, r[ix['任务子类型']])
    if jt == '未匹配分区':
        unmatched[str(part).strip()] += 1
    recs.append({
        '日期': t0.date(),
        '小时': t0.hour,
        '作业类型': jt,
        '拣货分区': str(part or '').strip(),
        '分区简码': zone_code(part),
        '拣货人': str(person).strip(),
        '行数': num(r[ix['拣货行数']]),
        '时长': hours,
        '数量': num(r[ix['拣货数量']]),
    })

print('有效记录: %d  |  丢弃: 无拣货人=%d, 时间缺失/异常=%d' % (len(recs), skip_empty_person, skip_no_time))
print('未匹配前后场分区: %d 条 %s' % (sum(unmatched.values()), dict(list(unmatched.items())[:8])))
print('日期范围: %s ~ %s' % (min(x['日期'] for x in recs), max(x['日期'] for x in recs)))


def agg(rs):
    h = sum(x['行数'] for x in rs)
    t = sum(x['时长'] for x in rs)
    return h, t, (h / t if t else float('nan'))


def table(title, keyfunc, top=None):
    print('\n===== %s =====' % title)
    print('%-28s %9s %10s %13s' % ('分组', '行数', '时长(h)', '效率(行/h)'))
    g = OrderedDict()
    for x in recs:
        g.setdefault(keyfunc(x), []).append(x)
    items = sorted(g.items(), key=lambda kv: -agg(kv[1])[2])
    if top:
        items = items[:top]
    for k, rs in items:
        h, t, e = agg(rs)
        print('%-28s %9.0f %10.2f %13.2f' % (str(k), h, t, e))
    h, t, e = agg(recs)
    print('%-28s %9.0f %10.2f %13.2f' % ('总计', h, t, e))


table('按作业类型', lambda x: x['作业类型'])
table('按 作业类型 + 拣货分区（细分）', lambda x: '%s|%s' % (x['作业类型'], x['拣货分区']))
table('按日期', lambda x: x['日期'])
table('按拣货人', lambda x: x['拣货人'])
table('按 日期 + 作业类型 + 拣货人', lambda x: '%s|%s|%s' % (x['日期'], x['作业类型'], x['拣货人']))

# ---- 小时维度 ----
print('\n===== 按小时 =====')
print('%-8s %9s %10s %13s' % ('小时', '行数', '时长(h)', '效率(行/h)'))
gh = OrderedDict()
for x in recs:
    gh.setdefault(x['小时'], []).append(x)
for h in sorted(gh):
    hh, tt, e = agg(gh[h])
    print('%-8s %9.0f %10.2f %13.2f' % ('%d点' % h, hh, tt, e))
hh, tt, e = agg(recs)
print('%-8s %9.0f %10.2f %13.2f' % ('总计', hh, tt, e))

print('\n===== 按 小时 + 作业类型 =====')
print('%-20s %9s %10s %13s' % ('小时|作业类型', '行数', '时长(h)', '效率(行/h)'))
gh2 = OrderedDict()
for x in recs:
    gh2.setdefault((x['小时'], x['作业类型']), []).append(x)
for k in sorted(gh2):
    hh, tt, e = agg(gh2[k])
    print('%-20s %9.0f %10.2f %13.2f' % ('%d点|%s' % k, hh, tt, e))

# 小时透视：行=作业类型+拣货人, 列=小时, 值=效率(行/h)
print('\n===== 小时透视(行=作业类型+拣货人, 列=小时, 值=效率 行/h) =====')
hours = sorted({x['小时'] for x in recs})
grid = defaultdict(lambda: [0.0, 0.0])
for x in recs:
    grid[('%s|%s' % (x['作业类型'], x['拣货人']), x['小时'])][0] += x['行数']
    grid[('%s|%s' % (x['作业类型'], x['拣货人']), x['小时'])][1] += x['时长']
rowkeys = sorted({('%s|%s' % (x['作业类型'], x['拣货人'])) for x in recs})


def eff(h, t):
    return ('%.2f' % (h / t)) if t else '-'


lines = ['作业类型|拣货人,' + ','.join('%d点' % h for h in hours) + ',总计']
for k in rowkeys:
    th = tt = 0.0
    cells = []
    for h in hours:
        rh, rt = grid.get((k, h), [0.0, 0.0])
        th += rh
        tt += rt
        cells.append(eff(rh, rt))
    line = '%s,%s,%s' % (k, ','.join(cells), eff(th, tt))
    lines.append(line)
    print(line)
tot = ['合计']
for h in hours:
    rh = sum(grid.get((k, h), [0.0, 0.0])[0] for k in rowkeys)
    rt = sum(grid.get((k, h), [0.0, 0.0])[1] for k in rowkeys)
    tot.append(eff(rh, rt))
tot.append(eff(sum(x['行数'] for x in recs), sum(x['时长'] for x in recs)))
lines.append(','.join(tot))
open(r'h:\Hema-Picking-Efficiency\Hema-Picking-Efficiency\_eff_hour.txt', 'w',
     encoding='utf-8-sig').write('\n'.join(lines))
print('\n小时透视已写入 _eff_hour.txt')

# 细分（作业类型 + 拣货分区）x 小时
zk = []
for x in recs:
    k = '%s|%s' % (x['作业类型'], x['拣货分区'])
    if k not in zk:
        zk.append(k)
zlines = ['作业类型|拣货分区,' + ','.join('%d点' % h for h in hours) + ',全天']
print('\n===== 细分（作业类型+拣货分区）x 小时 效率(行/h) =====')
for k in zk:
    rs = [x for x in recs if '%s|%s' % (x['作业类型'], x['拣货分区']) == k]
    cells, th, tt = [], 0.0, 0.0
    for h in hours:
        sub = [x for x in rs if x['小时'] == h]
        if sub:
            _, _, e = agg(sub)
            a, b, _ = agg(sub)
            th += a
            tt += b
            cells.append('%.2f' % e)
        else:
            cells.append('-')
    cells.append('%.2f' % (th / tt if tt else 0))
    zlines.append('%s,%s' % (k, ','.join(cells)))
    print('%-34s %s' % (k, ' '.join(cells)))
open(r'h:\Hema-Picking-Efficiency\Hema-Picking-Efficiency\_eff_zone.txt', 'w', encoding='utf-8-sig').write('\n'.join(zlines))
print('\n细分已写入 _eff_zone.txt')

# 明细落地
with open(r'h:\Hema-Picking-Efficiency\Hema-Picking-Efficiency\_eff_export.txt', 'w', encoding='utf-8-sig') as f:
    f.write('日期,作业类型,拣货人,拣货行数,时长(h),效率(行/h)\n')
    g = OrderedDict()
    for x in recs:
        g.setdefault((x['日期'], x['作业类型'], x['拣货人']), []).append(x)
    for (d, jt, p), rs in sorted(g.items()):
        h, t, e = agg(rs)
        f.write('%s,%s,%s,%.0f,%.4f,%.2f\n' % (d, jt, p, h, t, e))
print('\n明细已写入 _eff_export.txt')
