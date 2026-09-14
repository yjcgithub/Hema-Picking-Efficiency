"""由 export 导出文件生成网站的内置示例数据 web/assets/js/data.js
口径与报表一致：效率 = Σ拣货行数 / Σ拣货时长(h)，作业类型由前后场映射生成。
"""
import datetime as dt
import json
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORT = os.path.join(os.path.dirname(ROOT), '门店实时效率监控报表V1.0.xlsx')
EXPORT = os.path.join(os.path.dirname(ROOT), 'export-1785428943043.xlsx')
OUT = os.path.join(ROOT, 'assets', 'js', 'data.js')

# 前后场映射（与报表'前后场填写'一致）
mapping = {}
for r in openpyxl.load_workbook(REPORT, data_only=True)['前后场填写'].iter_rows(min_row=2, values_only=True):
    if r and r[0] is not None:
        mapping[str(r[0]).strip()] = str(r[1]).strip() if r[1] else ''

ws = openpyxl.load_workbook(EXPORT, data_only=True)['data']
data = list(ws.iter_rows(values_only=True))
hdr = list(data[0])
ix = {}
for i, h in enumerate(hdr):
    if str(h).strip() not in ix:
        ix[str(h).strip()] = i


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def as_dt(v):
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, str) and v.strip():
        for f in ('%Y-%m-%d %H:%M:%S', '%Y/%m/%d %H:%M:%S'):
            try:
                return dt.datetime.strptime(v.strip(), f)
            except ValueError:
                pass
    return None


def job_type(part, sub):
    if sub == '拆零拣打一体':
        return '一体化'
    p = mapping.get(str(part).strip())
    return p or '未匹配分区'


def zone_code(z):
    s = str(z or '').strip()
    i = s.find(' ')
    return s[:i] if i > 0 else s


recs = []
dropped = 0
for r in data[1:]:
    if r[ix['拣货单号']] is None or not r[ix['拣货人']]:
        dropped += 1
        continue
    t0, t1 = as_dt(r[ix['拣货开始时间']]), as_dt(r[ix['拣货完成时间']])
    if not t0 or not t1:
        dropped += 1
        continue
    hrs = (t1 - t0).total_seconds() / 3600.0
    if hrs <= 0:
        dropped += 1
        continue
    recs.append({'date': t0.strftime('%Y-%m-%d'), 'hour': t0.hour,
                 'slot': t0.hour + (0.5 if t0.minute >= 30 else 0),
                 'jobType': job_type(r[ix['拣货分区']], r[ix['任务子类型']]),
                 'zone': str(r[ix['拣货分区']] or '').strip(),
                 'code': zone_code(r[ix['拣货分区']]),
                 'person': str(r[ix['拣货人']]).strip(),
                 'rows': num(r[ix['拣货行数']]), 'hours': hrs})


def eff(h, t):
    return round(h / t, 2) if t else None


def agg(rs):
    return sum(x['rows'] for x in rs), sum(x['hours'] for x in rs)


def grouped(key):
    g = {}
    for x in recs:
        g.setdefault(x[key], []).append(x)
    return g


hours = sorted({x['hour'] for x in recs})
dates = sorted({x['date'] for x in recs})
H, T = agg(recs)

by_jt = []
for k, rs in grouped('jobType').items():
    h, t = agg(rs)
    by_jt.append({'name': k, 'rows': round(h), 'hours': round(t, 4), 'eff': eff(h, t)})
by_jt.sort(key=lambda d: -d['eff'])

by_person = []
for k, rs in grouped('person').items():
    h, t = agg(rs)
    by_person.append({'name': k, 'rows': round(h), 'hours': round(t, 4), 'eff': eff(h, t)})
by_person.sort(key=lambda d: -d['eff'])

by_hour = []
gh = grouped('hour')
for h in hours:
    a, b = agg(gh[h])
    by_hour.append({'hour': h, 'rows': round(a), 'hours': round(b, 4), 'eff': eff(a, b)})

# 半小时刻度（仅供「作业类型 × 小时 效率」「各小时效率趋势」两张图；透视仍按整点小时）
slots = sorted({x['slot'] for x in recs})
gs = grouped('slot')
by_slot = []
for s in slots:
    a, b = agg(gs[s])
    by_slot.append({'slot': s, 'rows': round(a), 'hours': round(b, 4), 'eff': eff(a, b)})

# 细分：作业类型 x 拣货分区
type_rank = {d['name']: i for i, d in enumerate(by_jt)}
zone_acc = {}
for x in recs:
    zone_acc.setdefault((x['jobType'], x['zone']), []).append(x)
by_zone = []
for (typ, z), rs in zone_acc.items():
    a, b = agg(rs)
    by_zone.append({'type': typ, 'zone': z, 'code': zone_code(z), 'rows': round(a),
                    'hours': round(b, 4), 'eff': eff(a, b),
                    'share': round(a / (H or 1) * 1000) / 10})
by_zone.sort(key=lambda d: (type_rank[d['type']], -d['eff']))

zone_by_hour = {'hours': hours, 'series': []}
for z in by_zone:
    rs = [x for x in recs if x['jobType'] == z['type'] and x['zone'] == z['zone']]
    dat = []
    for h in hours:
        sub = [x for x in rs if x['hour'] == h]
        dat.append(eff(*agg(sub)) if sub else None)
    zone_by_hour['series'].append({'name': z['type'] + '·' + z['code'], 'type': z['type'], 'data': dat})

# 作业类型 x 小时
series_jt = []
for jt in [d['name'] for d in by_jt]:
    dat = []
    for h in hours:
        rs = [x for x in recs if x['jobType'] == jt and x['hour'] == h]
        dat.append(eff(*agg(rs)) if rs else None)
    series_jt.append({'name': jt, 'data': dat})

# 作业类型 x 半小时
series_jt_slot = []
for jt in [d['name'] for d in by_jt]:
    dat = []
    for s in slots:
        rs = [x for x in recs if x['jobType'] == jt and x['slot'] == s]
        dat.append(eff(*agg(rs)) if rs else None)
    series_jt_slot.append({'name': jt, 'data': dat})

# 分块人员 x 半小时（供「人均」口径在半小时刻度上取数）
groups_slot = []
for jt in [d['name'] for d in by_jt]:
    rs = [x for x in recs if x['jobType'] == jt]
    persons = []
    for p in sorted({x['person'] for x in rs}):
        rs0 = [x for x in rs if x['person'] == p]
        dat = []
        for s in slots:
            sub = [x for x in rs0 if x['slot'] == s]
            dat.append(eff(*agg(sub)) if sub else None)
        persons.append({'person': p, 'data': dat})
    groups_slot.append({'type': jt, 'persons': persons})

# 人员(含作业类型) x 小时
person_rows = []
for jt in [d['name'] for d in by_jt]:
    for p in sorted({x['person'] for x in recs if x['jobType'] == jt}):
        rs0 = [x for x in recs if x['jobType'] == jt and x['person'] == p]
        dat = []
        for h in hours:
            rs = [x for x in rs0 if x['hour'] == h]
            dat.append(eff(*agg(rs)) if rs else None)
        a, b = agg(rs0)
        person_rows.append({'jobType': jt, 'person': p, 'data': dat, 'rows': round(a),
                            'hours': round(b, 4), 'total': eff(a, b)})
person_rows.sort(key=lambda d: -d['total'])

# 按作业类型分组的人员结构（组头带各小时/整体小计）
groups = []
for jt in by_jt:
    rs = [x for x in recs if x['jobType'] == jt['name']]
    a, b = agg(rs)
    hourly = []
    for h in hours:
        sub = [x for x in rs if x['hour'] == h]
        hourly.append(eff(*agg(sub)) if sub else None)
    groups.append({'type': jt['name'], 'rows': round(a), 'hours': round(b, 4),
                   'total': eff(a, b), 'hourly': hourly,
                   'persons': [r for r in person_rows if r['jobType'] == jt['name']]})

dataset = {
    'meta': {'sourceFile': os.path.basename(EXPORT), 'dates': dates, 'hours': hours,
             'recordCount': len(recs), 'dropped': dropped},
    'totals': {'rows': round(H), 'hours': round(T, 4), 'eff': eff(H, T),
               'persons': len(by_person), 'jobTypes': len(by_jt)},
    'byJobType': by_jt,
    'byPerson': by_person,
    'byHour': by_hour,
    'byZone': by_zone,
    'jobTypeByHour': {'hours': hours, 'series': series_jt,
                      'total': [d['eff'] for d in by_hour]},
    'zoneByHour': zone_by_hour,
    'bySlot': by_slot,
    'jobTypeBySlot': {'slots': slots, 'series': series_jt_slot,
                      'total': [d['eff'] for d in by_slot], 'groups': groups_slot},
    'personByHour': {'hours': hours, 'rows': person_rows, 'groups': groups},
}

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, 'w', encoding='utf-8') as f:
    f.write('/* 由 tools/gen_data.py 生成，勿手改 */\n')
    f.write('window.HEMA_DATA = ')
    json.dump(dataset, f, ensure_ascii=False, separators=(',', ':'))
    f.write(';\n')

print('已生成:', OUT)
print('记录:', len(recs), '| 人:', len(by_person), '| 小时:', hours, '| 日期:', dates)
